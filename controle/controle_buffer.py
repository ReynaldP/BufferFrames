# -*- coding: utf-8 -*-
"""
Contrôle de validité « Buffer AJU Frames » (section 5 du cahier des charges).

Reproduit en local, sur les copies .xlsx, la logique de l'appli Apps Script :
  - lecture ZZ84 (colonnes repérées par en-tête, filtrage des en-têtes répétés),
  - statut par poste via DATA Poste (secours codé en dur),
  - agrégation par bloc via le mapping (réf standard + POSMOD + nouvelles réfs),
  - capacité L1 (Buffer = L1 FRAME, doublons LR COMMUN / LR NEO gérés),
  - statut ROUGE / VERT / PLEIN / DÉBORD et débord.

Usage :
  python controle_buffer.py [--dir DOSSIER] [--json SORTIE.json]
  --json exporte les valeurs brutes des onglets (pour le banc de test Node).
"""
import argparse
import datetime as dt
import json
import math
import os
import re
import shutil
import sys
import tempfile
import unicodedata
import warnings

import openpyxl

warnings.filterwarnings("ignore")

PILOTAGE = "Pilotage Ferrure_PC FRAME.xlsx"
DIMENSIONNEMENT = "Dimensionnement Buffer ZINC.xlsx"

FALLBACK_POSTES = {
    "37SBEL1": "BUFFER",
    "53B0223": "AJU", "53B0225": "AJU", "37S0223": "AJU", "37S0225": "AJU",
}

# (prog, version, bloc, flux, réf L1, [(réf, type, version réf, à valider)])
MAPPING = [
    ("A320", "NEO/CEO", "C36 AV G", "Amboise", "D571.30067.206.02", [("D571.32444.200.00", "POSMOD", "", False)]),
    ("A320", "NEO/CEO", "C36 AV D", "Amboise", "D571.30067.207.02", [("D571.32444.201.00", "POSMOD", "", False)]),
    ("A320", "NEO/CEO", "C36 AR G", "M5", "D571.30068.200.02", [("D571.32445.200.00", "POSMOD", "", False)]),
    ("A320", "NEO/CEO", "C36 AR D", "M5", "D571.30068.201.02", [("D571.32445.201.00", "POSMOD", "", False)]),
    ("A320", "NEO/CEO", "C42 AV G", "Amboise", "D571.30287.200.02", [("D571.32209.200.00", "NOUVELLE", "", False),
                                                                     ("D571.30422.200.00", "CJ", "CJ", True)]),
    ("A320", "NEO/CEO", "C42 AV D", "Amboise", "D571.30287.201.03", [("D571.32209.202.00", "NOUVELLE", "", False)]),
    ("A320", "NEO/CEO", "C42 AR G", "T5X→H5", "D571.30410.200.00", [("D571.32210.200.40", "NOUVELLE", "", True),
                                                                    ("E571.11161.200.00", "NOUVELLE", "", True)]),
    ("A320", "NEO/CEO", "C42 AR D", "T5X→H5", "D571.30410.201.00", [("D571.32210.202.40", "NOUVELLE", "", True),
                                                                    ("E571.11161.202.00", "NOUVELLE", "", True)]),
    ("A321", "ACF", "C36 AV G", "Amboise", "D571.34310.200.01", [("E571.11135.200.00", "POSMOD", "", False)]),
    ("A321", "ACF", "C36 AV D", "Amboise", "D571.34310.201.01", [("E571.11135.201.00", "POSMOD", "", False)]),
    ("A321", "ACF", "C36 AR G", "M5", "D571.34311.200.01", [("E571.11136.200.00", "POSMOD", "", False)]),
    ("A321", "ACF", "C36 AR D", "M5", "D571.34311.201.01", [("E571.11136.201.00", "POSMOD", "", False)]),
    ("A321", "ACF", "C42 AV G", "M5", "D571.34312.200.01", []),
    ("A321", "ACF", "C42 AV D", "M5", "D571.34313.200.01", []),
    ("A321", "ACF", "C42 AR G", "T5X→H5", "D571.34314.200.01", [("E571.11163.200.01", "NOUVELLE", "", True)]),
    ("A321", "ACF", "C42 AR D", "T5X→H5", "D571.34319.200.01", [("E571.11165.200.01", "NOUVELLE", "", True)]),
    ("A321", "XLR", "C42 AV G", "M5", "D571.34812.202.00", []),
    ("A321", "XLR", "C42 AV D", "M5", "D571.34813.202.00", []),
    ("A321", "XLR", "C42 AR G", "T5X→H5", "D571.34814.200.00", []),
    ("A321", "XLR", "C42 AR D", "T5X→H5", "D571.34819.200.00", []),
    ("A330", "NEO", "C40 AV G", "M5", "F571.16770.200.00", []),
    ("A330", "NEO", "C40 AV D", "M5", "F571.16770.201.00", []),
    ("A330", "CEO", "C40 AV G", "M5", "F571.12887.200.01", []),
    ("A330", "CEO", "C40 AV D", "M5", "F571.12887.201.01", []),
    ("A330", "COMMUN", "C40 AR G", "V4", "F571.12262.202.01", []),
    ("A330", "COMMUN", "C40 AR D", "V4", "F571.12262.203.01", []),
    ("A330", "NEO/CEO", "C47 G", "M5", "F571.12994.200.00", []),
    ("A330", "NEO/CEO", "C47 D", "M5", "F571.12994.201.00", []),
]

# Résultats attendus (section 5) : (prog, version, bloc) -> (AJU, À venir)
EXPECTED = {
    ("A320", "NEO/CEO", "C36 AV G"): (3, 9), ("A320", "NEO/CEO", "C36 AV D"): (9, 0),
    ("A320", "NEO/CEO", "C36 AR G"): (4, 0), ("A320", "NEO/CEO", "C36 AR D"): (4, 0),
    ("A320", "NEO/CEO", "C42 AV G"): (1, 4), ("A320", "NEO/CEO", "C42 AV D"): (0, 4),
    ("A320", "NEO/CEO", "C42 AR G"): (0, 2), ("A320", "NEO/CEO", "C42 AR D"): (1, 2),
    ("A321", "ACF", "C36 AV G"): (9, 4), ("A321", "ACF", "C36 AV D"): (8, 6),
    ("A321", "ACF", "C36 AR G"): (12, 8), ("A321", "ACF", "C36 AR D"): (12, 8),
    ("A321", "ACF", "C42 AV G"): (6, 0), ("A321", "ACF", "C42 AV D"): (4, 0),
    ("A321", "ACF", "C42 AR G"): (3, 4), ("A321", "ACF", "C42 AR D"): (13, 4),
    ("A321", "XLR", "C42 AV G"): (10, 0), ("A321", "XLR", "C42 AV D"): (6, 0),
    ("A321", "XLR", "C42 AR G"): (0, 1), ("A321", "XLR", "C42 AR D"): (0, 1),
}
EXPECTED_A330_AJU = {("A330", "COMMUN", "C40 AR G"): 2, ("A330", "NEO", "C40 AV G"): 2, ("A330", "NEO", "C40 AV D"): 2}

YEAR_INDEX = {"2024": 9, "2025": 10, "2026-2028": 11, "2028+": 12}


def norm_ref(v):
    return re.sub(r"\s+", "", str(v if v is not None else "")).upper()


def norm_head(v):
    s = unicodedata.normalize("NFD", str(v if v is not None else "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def to_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return 0.0


def open_wb(path):
    """Ouvre un classeur ; s'il est verrouillé (Excel / OneDrive), travaille sur une copie."""
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        tmp = os.path.join(tempfile.mkdtemp(), os.path.basename(path))
        shutil.copy2(path, tmp)
        return openpyxl.load_workbook(tmp, data_only=True)


def sheet_values(wb, name):
    ws = wb[name]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def read_postes(rows):
    m = {}
    for r in rows:
        if len(r) < 4:
            continue
        poste, statut = norm_ref(r[2]), norm_ref(r[3])
        if poste and statut in ("BUFFER", "AJU"):
            m[poste] = statut
    if "BUFFER" not in m.values() or "AJU" not in m.values():
        for k, v in FALLBACK_POSTES.items():
            m.setdefault(k, v)
    return m


ZZ84_COLS = {"oper": ["oper"], "ordre": ["ordre"], "article": ["article"], "designation": ["designation"],
             "poste": ["postrav", "postetravail", "poste"], "qte": ["qteof", "qte"], "age": ["agepos", "age"],
             "serie": ["nserie", "noserie", "numeroserie", "serie"]}
ZZ84_FALLBACK = {"oper": 2, "ordre": 3, "article": 4, "designation": 9, "poste": 12, "qte": 13, "age": 15, "serie": 16}


def read_zz84(rows, postes):
    header_row, cols = None, {}
    for i, r in enumerate(rows[:15]):
        heads = [norm_head(v) for v in r]
        if "article" in heads and any(h in heads for h in ("postrav", "postetravail")):
            header_row = i
            for key, names in ZZ84_COLS.items():
                for n in names:
                    if n in heads:
                        cols[key] = heads.index(n)
                        break
            break
    if header_row is None:
        header_row = 2
    for k, v in ZZ84_FALLBACK.items():
        cols.setdefault(k, v)
    pieces, skipped_headers = [], 0
    for r in rows[header_row + 1:]:
        get = lambda k: r[cols[k]] if cols[k] < len(r) else None
        ref, poste = norm_ref(get("article")), norm_ref(get("poste"))
        if not ref or not poste:
            continue
        if poste in ("POS.TRAV", "POSTRAV") or ref == "ARTICLE":
            skipped_headers += 1
            continue
        statut = postes.get(poste)
        if not statut:
            continue
        pieces.append({"ref": ref, "statut": statut, "qte": to_num(get("qte")), "age": to_num(get("age")),
                       "of": get("ordre"), "des": str(get("designation") or "").strip(), "poste": poste})
    return pieces, skipped_headers, rows[0][5], rows[0][6]


def read_l1(rows, year):
    idx = YEAR_INDEX[year]
    caps, dups = {}, []
    for r in rows[1:]:
        if not r or not r[0] or "FRAME" not in str(r[5] or "").upper():
            continue
        ref = norm_ref(r[0])
        assemblage = str(r[2] or "").upper()
        if ref in caps:
            dups.append((ref, assemblage))
            if "COMMUN" in assemblage and "COMMUN" not in caps[ref][1]:
                caps[ref] = (to_num(r[idx]), assemblage)
            continue
        caps[ref] = (to_num(r[idx]), assemblage)
    return {k: v[0] for k, v in caps.items()}, dups


def statut_bloc(aju, mini, maxi, cap):
    if aju > cap:
        return "DÉBORD"
    if aju < mini:
        return "ROUGE"
    if aju < maxi:
        return "VERT"
    return "PLEIN"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    ap.add_argument("--year", default="2026-2028")
    ap.add_argument("--json", help="exporte les valeurs brutes des onglets en JSON")
    args = ap.parse_args()

    wb_p = open_wb(os.path.join(args.dir, PILOTAGE))
    wb_d = open_wb(os.path.join(args.dir, DIMENSIONNEMENT))
    zz84_rows = sheet_values(wb_p, "ZZ84")
    poste_rows = sheet_values(wb_p, "DATA Poste")
    l1_rows = sheet_values(wb_d, "L1")

    if args.json:
        def enc(v):
            if isinstance(v, dt.datetime):
                return {"__date": v.isoformat()}
            if isinstance(v, dt.time):
                return {"__date": dt.datetime(1899, 12, 30, v.hour, v.minute, v.second).isoformat()}
            return "" if v is None else v
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump({"ZZ84": [[enc(v) for v in r] for r in zz84_rows],
                       "DATA Poste": [[enc(v) for v in r] for r in poste_rows],
                       "L1": [[enc(v) for v in r] for r in l1_rows]}, f, ensure_ascii=False)

    postes = read_postes(poste_rows)
    pieces, skipped, d_extr, h_extr = read_zz84(zz84_rows, postes)
    caps, dups = read_l1(l1_rows, args.year)

    ref2bloc, blocs = {}, {}
    for prog, ver, bloc, flux, ref_l1, others in MAPPING:
        key = (prog, ver, bloc)
        blocs[key] = {"flux": flux, "ref_l1": ref_l1, "aju": 0.0, "avenir": 0.0, "age_max": 0.0}
        ref2bloc[norm_ref(ref_l1)] = key
        for ref, *_ in others:
            ref2bloc[norm_ref(ref)] = key

    unmapped = {}
    for p in pieces:
        key = ref2bloc.get(p["ref"])
        if key is None:
            unmapped[p["ref"]] = unmapped.get(p["ref"], 0) + p["qte"]
            continue
        b = blocs[key]
        if p["statut"] == "AJU":
            b["aju"] += p["qte"]
            b["age_max"] = max(b["age_max"], p["age"])
        else:
            b["avenir"] += p["qte"]

    print("Buffer AJU Frames : contrôle de validité")
    print("Extraction SAP : %s à %s" % (d_extr.strftime("%d/%m/%Y") if hasattr(d_extr, "strftime") else d_extr,
                                         h_extr.strftime("%H:%M") if hasattr(h_extr, "strftime") else h_extr))
    print("Postes BUFFER : %s" % ", ".join(sorted(k for k, v in postes.items() if v == "BUFFER")))
    print("Postes AJU    : %s" % ", ".join(sorted(k for k, v in postes.items() if v == "AJU")))
    print("Lignes d'en-tête répétées filtrées : %d" % skipped)
    print("Doublons L1 ignorés : %s" % ", ".join("%s (%s)" % d for d in dups))
    print("Année L1 : %s\n" % args.year)

    fmt = "%-4s %-8s %-9s %-8s %4s %4s %5s %5s %5s %5s  %-7s %s"
    print(fmt % ("Prog", "Version", "Bloc", "Flux", "AJU", "Ven", "Cap", "Mini", "Maxi", "Déb", "Statut", "Contrôle"))
    ok_all = True
    tot = {"A320+A321": [0, 0], "A330": [0, 0]}
    debord = 0
    for key, b in blocs.items():
        cap = caps.get(norm_ref(b["ref_l1"]), 0)
        mini = max(1, math.ceil(0.25 * cap)) if cap > 0 else 0
        maxi = cap
        st = statut_bloc(b["aju"], mini, maxi, cap)
        exc = max(0, b["aju"] - cap)
        debord += exc
        check = ""
        if key in EXPECTED:
            exp = EXPECTED[key]
            good = (b["aju"], b["avenir"]) == exp
            check = "OK" if good else "ÉCART (attendu %d / %d)" % exp
            ok_all &= good
        elif key in EXPECTED_A330_AJU or key[0] == "A330":
            exp = EXPECTED_A330_AJU.get(key, 0)
            good = b["aju"] == exp
            check = "OK" if good else "ÉCART (attendu AJU %d)" % exp
            ok_all &= good
        grp = "A330" if key[0] == "A330" else "A320+A321"
        tot[grp][0] += b["aju"]
        tot[grp][1] += b["avenir"]
        print(fmt % (key[0], key[1], key[2], b["flux"], int(b["aju"]), int(b["avenir"]), int(cap), mini, int(maxi),
                     int(exc), st, check))

    print("\nTotaux A320 + A321 : AJU = %d, À venir = %d (attendu 105 / 57)" % tuple(tot["A320+A321"]))
    print("Totaux A330        : AJU = %d, À venir = %d (attendu AJU 6)" % tuple(tot["A330"]))
    print("(Débord si chaque variante avait son propre rack : %d)" % debord)

    # Racks physiques : les variantes d'un même bloc (A321 ACF + XLR, A330 NEO + CEO) partagent la place
    racks = {}
    for key, b in blocs.items():
        cap = caps.get(norm_ref(b["ref_l1"]), 0)
        mini = max(1, math.ceil(0.25 * cap)) if cap > 0 else 0
        r = racks.setdefault((key[0], key[2]), {"versions": [], "aju": 0, "avenir": 0, "cap": 0, "mini": 0, "maxi": 0})
        r["versions"].append(key[1])
        for f, v in (("aju", b["aju"]), ("avenir", b["avenir"]), ("cap", cap), ("mini", mini), ("maxi", cap)):
            r[f] += v
    print("\nRacks partagés (capacité, mini et maxi additionnés) :")
    rack_debord = 0
    for (prog, bloc), r in racks.items():
        exc = max(0, r["aju"] - r["cap"])
        rack_debord += exc
        if len(r["versions"]) > 1 or exc:
            print("  %-4s %-9s %-14s AJU %3d / %3d  mini %2d  maxi %2d  %-7s débord %d" % (
                prog, bloc, " + ".join(r["versions"]), r["aju"], r["cap"], r["mini"], r["maxi"],
                statut_bloc(r["aju"], r["mini"], r["maxi"], r["cap"]), exc))
    print("Débord total (racks partagés) : %d" % rack_debord)
    print("Réfs BUFFER/AJU non rattachées : %s" % (unmapped or "aucune"))
    ok_all &= tot["A320+A321"] == [105, 57] and tot["A330"][0] == 6 and not unmapped
    print("\nRÉSULTAT : %s" % ("CONFORME" if ok_all else "NON CONFORME"))
    return 0 if ok_all else 1


if __name__ == "__main__":
    sys.exit(main())
